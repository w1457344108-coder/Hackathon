"""Source document discovery and loading helpers.

Raw source documents are read-only. This module extracts plain text into
`processed/texts/` and writes metadata records without modifying originals.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Iterable

from config import (
    DEFAULT_REVIEW_STATUS,
    PROCESSED_METADATA_DIR,
    PROCESSED_TEXTS_DIR,
    PROJECT_ROOT,
    RAW_DOCUMENTS_DIR,
    SUPPORTED_EXTENSIONS,
)
from modules.text_cleaner import clean_text


LogFn = Callable[[str], None]


@dataclass
class DocumentMetadata:
    doc_id: str
    original_filename: str
    file_type: str
    text_file_path: str
    character_count: int
    word_count: int
    read_status: str
    error_message: str
    created_at: str


def discover_documents(documents_dir: Path = RAW_DOCUMENTS_DIR) -> list[Path]:
    """Return supported, non-hidden files in deterministic filename order."""

    if not documents_dir.exists():
        return []

    return sorted(
        (
            path
            for path in documents_dir.iterdir()
            if path.is_file()
            and not path.name.startswith(".")
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        ),
        key=lambda item: item.name.lower(),
    )


def discover_all_source_files(documents_dir: Path = RAW_DOCUMENTS_DIR) -> list[Path]:
    """Return all non-hidden files so unsupported files can be reported."""

    if not documents_dir.exists():
        return []

    return sorted(
        (path for path in documents_dir.iterdir() if path.is_file() and not path.name.startswith(".")),
        key=lambda item: item.name.lower(),
    )


def create_doc_id(path: Path) -> str:
    """Create a stable filesystem-safe id from a filename."""

    stem = path.stem.lower()
    slug = re.sub(r"[^a-z0-9]+", "_", stem).strip("_")
    if not slug:
        slug = "document"
    digest = hashlib.sha1(path.name.encode("utf-8")).hexdigest()[:10]
    return f"{slug}_{digest}"


def extract_text_from_file(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix == ".docx":
        return extract_docx_text(path)
    if suffix in {".txt", ".md"}:
        return extract_plain_text(path)
    if suffix in {".html", ".htm"}:
        return extract_html_text(path)

    raise ValueError(f"Unsupported file type: {suffix}")


def extract_pdf_text(path: Path) -> str:
    try:
        import fitz  # type: ignore
    except ImportError as exc:
        raise RuntimeError("PyMuPDF is required to read PDF files. Install pymupdf.") from exc

    parts: list[str] = []
    with fitz.open(path) as document:
        for page_index, page in enumerate(document, start=1):
            page_text = page.get_text("text")
            if page_text.strip():
                parts.append(f"[Page {page_index}]\n{page_text}")
    return "\n\n".join(parts)


def extract_docx_text(path: Path) -> str:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise RuntimeError("python-docx is required to read DOCX files. Install python-docx.") from exc

    document = Document(path)
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def extract_plain_text(path: Path) -> str:
    errors: list[str] = []
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: {exc}")

    raise UnicodeDecodeError(
        "unknown",
        b"",
        0,
        1,
        "Unable to decode text file with utf-8-sig, utf-8, gb18030, or latin-1. "
        + " | ".join(errors),
    )


def extract_html_text(path: Path) -> str:
    try:
        from bs4 import BeautifulSoup  # type: ignore
    except ImportError as exc:
        raise RuntimeError("beautifulsoup4 is required to read HTML files. Install beautifulsoup4.") from exc

    html = extract_plain_text(path)
    soup = BeautifulSoup(html, "html.parser")
    for item in soup(["script", "style"]):
        item.decompose()
    return soup.get_text("\n")


def count_words(text: str) -> int:
    return len(re.findall(r"[\w]+", text, flags=re.UNICODE))


def relative_path(path: Path) -> str:
    try:
        return path.relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        return path.as_posix()


def write_metadata(metadata: Iterable[DocumentMetadata], metadata_dir: Path = PROCESSED_METADATA_DIR) -> None:
    metadata_dir.mkdir(parents=True, exist_ok=True)
    records = [asdict(item) for item in metadata]

    json_path = metadata_dir / "documents_metadata.json"
    csv_path = metadata_dir / "documents_metadata.csv"

    json_path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")

    fieldnames = [
        "doc_id",
        "original_filename",
        "file_type",
        "text_file_path",
        "character_count",
        "word_count",
        "read_status",
        "error_message",
        "created_at",
    ]
    with csv_path.open("w", newline="", encoding="utf-8-sig") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)


def process_documents(
    documents_dir: Path = RAW_DOCUMENTS_DIR,
    texts_dir: Path = PROCESSED_TEXTS_DIR,
    metadata_dir: Path = PROCESSED_METADATA_DIR,
    log: LogFn = print,
) -> list[DocumentMetadata]:
    """Extract and clean all source documents, recording failures as metadata."""

    texts_dir.mkdir(parents=True, exist_ok=True)
    metadata_dir.mkdir(parents=True, exist_ok=True)

    source_files = discover_all_source_files(documents_dir)
    metadata: list[DocumentMetadata] = []

    log(f"扫描原始文献目录: {documents_dir}")
    log(f"发现文件数量: {len(source_files)}")

    for source_file in source_files:
        doc_id = create_doc_id(source_file)
        text_path = texts_dir / f"{doc_id}.txt"
        created_at = datetime.now(timezone.utc).isoformat()
        file_type = source_file.suffix.lower()

        log(f"正在读取: {source_file.name}")

        if file_type not in SUPPORTED_EXTENSIONS:
            message = f"Unsupported file type: {file_type}"
            log(f"读取失败: {source_file.name} - {message}")
            metadata.append(
                DocumentMetadata(
                    doc_id=doc_id,
                    original_filename=source_file.name,
                    file_type=file_type,
                    text_file_path="",
                    character_count=0,
                    word_count=0,
                    read_status="failed",
                    error_message=message,
                    created_at=created_at,
                )
            )
            continue

        try:
            extracted = extract_text_from_file(source_file)
            cleaned = clean_text(extracted)

            if not cleaned:
                raise ValueError("No readable text was extracted.")

            text_path.write_text(cleaned, encoding="utf-8")
            log(f"读取成功: {source_file.name}")
            log(f"文本已保存: {text_path}")

            metadata.append(
                DocumentMetadata(
                    doc_id=doc_id,
                    original_filename=source_file.name,
                    file_type=file_type,
                    text_file_path=relative_path(text_path),
                    character_count=len(cleaned),
                    word_count=count_words(cleaned),
                    read_status="success",
                    error_message="",
                    created_at=created_at,
                )
            )
        except Exception as exc:
            message = str(exc)
            log(f"读取失败: {source_file.name} - {message}")
            metadata.append(
                DocumentMetadata(
                    doc_id=doc_id,
                    original_filename=source_file.name,
                    file_type=file_type,
                    text_file_path="",
                    character_count=0,
                    word_count=0,
                    read_status="failed",
                    error_message=message,
                    created_at=created_at,
                )
            )

    write_metadata(metadata, metadata_dir)
    log(f"元数据 CSV 已保存: {metadata_dir / 'documents_metadata.csv'}")
    log(f"元数据 JSON 已保存: {metadata_dir / 'documents_metadata.json'}")

    success_count = sum(1 for item in metadata if item.read_status == "success")
    failed_count = len(metadata) - success_count
    log(f"文献读取完成: success={success_count}, failed={failed_count}")

    # Keep the imported constant visible to future generated records and linters.
    _ = DEFAULT_REVIEW_STATUS
    return metadata
