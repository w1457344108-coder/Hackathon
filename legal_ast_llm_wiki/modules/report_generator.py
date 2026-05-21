"""Final report generation helpers for the AST + LLM Wiki prototype."""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from config import OUTPUT_REPORTS_DIR, PROJECT_ROOT
from modules.chunker import backup_existing_file


FINAL_REPORT_MD_NAME = "final_ast_llm_wiki_report.md"
FINAL_REPORT_DOCX_NAME = "final_ast_llm_wiki_report.docx"


def read_json_list(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError:
        return []
    return data if isinstance(data, list) else []


def parse_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return [text]
    if isinstance(parsed, list):
        return [str(item).strip() for item in parsed if str(item).strip()]
    return [text]


def count_markdown_pages(wiki_dir: Path) -> dict[str, int]:
    counts: dict[str, int] = {}
    if not wiki_dir.exists():
        return counts
    for path in wiki_dir.rglob("*.md"):
        if ".bak" in path.name or "_backups" in path.parts:
            continue
        relative = path.relative_to(wiki_dir)
        key = relative.parts[0] if len(relative.parts) > 1 else "root"
        counts[key] = counts.get(key, 0) + 1
    counts["total"] = sum(counts.values())
    return counts


def collect_project_metrics(project_root: Path = PROJECT_ROOT) -> dict[str, Any]:
    metadata = read_json_list(project_root / "processed/metadata/documents_metadata.json")
    chunks = read_json_list(project_root / "processed/chunks/document_chunks.json")
    ast_nodes = read_json_list(project_root / "processed/ast_nodes/ast_nodes.json")
    legal_terms = read_json_list(project_root / "outputs/json/legal_terms.json")
    colloquial_mappings = read_json_list(project_root / "outputs/json/colloquial_mapping.json")
    synonyms = read_json_list(project_root / "outputs/json/synonym_table.json")
    retrieval_results = read_json_list(project_root / "outputs/json/retrieval_results.json")
    wiki_counts = count_markdown_pages(project_root / "wiki")

    chunks_by_doc = Counter(str(chunk.get("doc_id", "")) for chunk in chunks if str(chunk.get("doc_id", "")).strip())
    ast_by_doc = Counter(str(node.get("doc_id", "")) for node in ast_nodes if str(node.get("doc_id", "")).strip())
    domains = Counter(str(node.get("legal_domain", "unknown")) or "unknown" for node in ast_nodes)
    term_categories = Counter(str(term.get("category", "unknown")) or "unknown" for term in legal_terms)

    retrieval_scores = [
        float(result.get("relevance_score", 0) or 0)
        for result in retrieval_results
        if str(result.get("relevance_score", "")).strip() != ""
    ]
    retrieval_with_terms = sum(1 for result in retrieval_results if parse_list(result.get("mapped_terms", [])))
    retrieval_with_ast = sum(1 for result in retrieval_results if parse_list(result.get("matched_ast_nodes", [])) or result.get("matched_ast_nodes"))
    retrieval_with_wiki = sum(1 for result in retrieval_results if parse_list(result.get("matched_wiki_pages", [])) or result.get("matched_wiki_pages"))
    retrieval_with_sources = sum(1 for result in retrieval_results if parse_list(result.get("source_doc_ids", [])))

    documents = []
    for record in metadata:
        doc_id = str(record.get("doc_id", ""))
        documents.append(
            {
                "doc_id": doc_id,
                "filename": record.get("original_filename", ""),
                "file_type": record.get("file_type", ""),
                "read_status": record.get("read_status", ""),
                "characters": record.get("character_count", 0),
                "words": record.get("word_count", 0),
                "chunks": chunks_by_doc.get(doc_id, 0),
                "ast_nodes": ast_by_doc.get(doc_id, 0),
            }
        )

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "document_count": len(metadata),
        "successful_documents": sum(1 for item in metadata if item.get("read_status") == "success"),
        "failed_documents": sum(1 for item in metadata if item.get("read_status") != "success"),
        "documents": documents,
        "chunk_count": len(chunks),
        "ast_node_count": len(ast_nodes),
        "wiki_page_count": wiki_counts.get("total", 0),
        "wiki_counts": wiki_counts,
        "legal_term_count": len(legal_terms),
        "colloquial_mapping_count": len(colloquial_mappings),
        "synonym_count": len(synonyms),
        "retrieval_query_count": len(retrieval_results),
        "retrieval_average_score": round(sum(retrieval_scores) / len(retrieval_scores), 4) if retrieval_scores else 0,
        "retrieval_with_terms": retrieval_with_terms,
        "retrieval_with_ast": retrieval_with_ast,
        "retrieval_with_wiki": retrieval_with_wiki,
        "retrieval_with_sources": retrieval_with_sources,
        "domain_distribution": domains.most_common(),
        "term_category_distribution": term_categories.most_common(),
        "retrieval_results": retrieval_results,
    }


def table_row(label: str, value: Any) -> str:
    return f"| {label} | {value} |"


def markdown_table(headers: list[str], rows: list[list[Any]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    for row in rows:
        clean = [re.sub(r"\s+", " ", str(cell)).replace("|", "/") for cell in row]
        lines.append("| " + " | ".join(clean) + " |")
    return "\n".join(lines)


def build_final_report_markdown(metrics: dict[str, Any]) -> str:
    documents = metrics.get("documents", [])
    retrieval_results = metrics.get("retrieval_results", [])

    doc_rows = [
        [
            item.get("doc_id", ""),
            item.get("filename", ""),
            item.get("file_type", ""),
            item.get("read_status", ""),
            item.get("chunks", 0),
            item.get("ast_nodes", 0),
        ]
        for item in documents
    ]
    wiki_rows = [
        [key, value]
        for key, value in sorted(metrics.get("wiki_counts", {}).items())
        if key != "total"
    ]
    domain_rows = metrics.get("domain_distribution", [])[:10]
    term_category_rows = metrics.get("term_category_distribution", [])[:12]
    retrieval_rows = [
        [
            result.get("query_id", ""),
            result.get("query_text", ""),
            ", ".join(parse_list(result.get("mapped_terms", []))[:4]),
            len(result.get("matched_ast_nodes", []) or []),
            len(result.get("matched_wiki_pages", []) or []),
            result.get("relevance_score", 0),
        ]
        for result in retrieval_results
    ]

    lines = [
        "# AST + LLM Wiki 法律知识库原型最终构建报告",
        "",
        f"生成时间：{metrics.get('generated_at', '')}",
        "",
        "## 1. 项目背景",
        "",
        "本项目围绕 9 篇新加坡相关法律、政策与数字治理文献，构建一个可运行的法律知识库原型。原型目标不是直接替代法律专业判断，而是验证能否将原始文献转化为可追溯、可维护、可检索的结构化知识资产。",
        "",
        "## 2. 构建目标",
        "",
        "项目目标包括：读取原始文献、转换纯文本、进行法律/政策文本切块、生成类 AST 结构化节点、构建 LLM Wiki、建立专业词汇库与口语表达映射，并通过轻量检索实验观察方案可行性。",
        "",
        "## 3. 数据来源说明",
        "",
        "原始文献统一放置在 `legal_ast_llm_wiki/raw/documents/`。流水线只读取原始文件，不修改、重命名或删除原始文献。当前样本数量有限，仅用于原型验证，不代表完整的新加坡法律知识库。",
        "",
        "## 4. 九篇文献处理情况",
        "",
        markdown_table(["doc_id", "文件名", "格式", "读取状态", "chunk 数", "AST 节点数"], doc_rows),
        "",
        "## 5. 系统架构说明",
        "",
        "系统采用模块化流水线：`file_loader` 负责读取文件，`text_cleaner` 负责保守清洗，`chunker` 负责法律/政策友好的切块，`ast_parser` 负责本地规则解析，`wiki_builder` 负责 Markdown Wiki，`glossary_builder` 负责词汇与映射，`retriever` 和 `evaluator` 负责检索验证，`report_generator` 负责最终报告。",
        "",
        "关键产物统计如下：",
        "",
        markdown_table(
            ["指标", "数量"],
            [
                ["文献数量", metrics.get("document_count", 0)],
                ["成功读取文献", metrics.get("successful_documents", 0)],
                ["失败读取文献", metrics.get("failed_documents", 0)],
                ["chunk 数量", metrics.get("chunk_count", 0)],
                ["AST 节点数量", metrics.get("ast_node_count", 0)],
                ["Wiki 页面数量", metrics.get("wiki_page_count", 0)],
                ["专业术语数量", metrics.get("legal_term_count", 0)],
                ["口语映射数量", metrics.get("colloquial_mapping_count", 0)],
                ["同义词/近义词记录", metrics.get("synonym_count", 0)],
                ["检索测试问题数量", metrics.get("retrieval_query_count", 0)],
            ],
        ),
        "",
        "Wiki 页面分布：",
        "",
        markdown_table(["Wiki 区域", "页面数"], wiki_rows),
        "",
        "## 6. 为什么采用“类 AST + LLM Wiki”",
        "",
        "普通全文检索主要依赖关键词或相似度，难以显式区分法律主体、行为、对象、条件、义务、例外和后果。类 AST 将法律/政策文本拆成可检索字段；LLM Wiki 则把这些字段组织为人类、Obsidian 和后续 Agent/LLM 都能阅读的知识层。二者结合的主要价值是提高检索精确性、可解释性、可追溯性和知识维护能力。",
        "",
        "## 7. 类 AST 节点设计",
        "",
        "本项目中的 AST 不是代码语法树，而是法律/政策文本的结构化解析树。每个节点保留 `doc_id`、`chunk_id`、`source_file`、`original_text`、`source_evidence` 等来源字段，并尝试抽取法域、文献类型、法律领域、法律主体、法律行为、法律对象、适用条件、义务、权利、禁止事项、法律后果、例外、定义、专业术语、相关术语、用户可能问题、置信度与审核状态。",
        "",
        "识别出的法律领域分布：",
        "",
        markdown_table(["法律领域", "节点数"], domain_rows),
        "",
        "## 8. LLM Wiki 页面结构",
        "",
        "LLM Wiki 不是回答模型，而是知识存储、组织和维护层。它包括来源文献页面、法律/政策主题页面、概念页面、问题页面、术语表页面和结构化节点页面。后续 Agent 或 LLM 可以基于 Wiki 进行问答、检索、核验和更新，但 Wiki 本身不直接提供最终法律结论。",
        "",
        "## 9. 专业词汇库构建方法",
        "",
        "词汇库从 AST 节点的法律主体、行为、对象、条件、义务、权利、禁止事项、法律后果、定义、专业术语和相关术语中抽取候选词，并保留来源节点、来源文献和证据片段。当前词汇库为候选结果，全部需要人工审核。",
        "",
        markdown_table(["词汇类别", "数量"], term_category_rows),
        "",
        "## 10. 口语表达映射方法",
        "",
        "口语表达映射用于把普通用户的问题表达映射到专业术语、法律问题或概念。例如，关于“send my data overseas”的问题会尝试映射到跨境传输或 personal data transfer 相关术语。当前映射由规则和 AST 证据生成，仍属于待审核候选结果。",
        "",
        "## 11. 检索验证方法",
        "",
        "检索验证使用 8 个英文问题，采用本地轻量方法：TF-IDF 相似度、关键词命中、专业术语匹配、AST 字段权重、Wiki 标题与 YAML tags 匹配，以及原文 chunk 相似度。检索只输出候选结果和证据，不冒充正式法律咨询。",
        "",
        "## 12. 检索验证结果",
        "",
        f"平均增强检索得分为 {metrics.get('retrieval_average_score', 0)}。{metrics.get('retrieval_with_terms', 0)}/{metrics.get('retrieval_query_count', 0)} 个问题返回了专业术语，{metrics.get('retrieval_with_ast', 0)}/{metrics.get('retrieval_query_count', 0)} 个问题返回了 AST 节点，{metrics.get('retrieval_with_wiki', 0)}/{metrics.get('retrieval_query_count', 0)} 个问题返回了 Wiki 页面，{metrics.get('retrieval_with_sources', 0)}/{metrics.get('retrieval_query_count', 0)} 个问题返回了来源文献。",
        "",
        markdown_table(["query_id", "问题", "映射术语示例", "AST 命中数", "Wiki 命中数", "得分"], retrieval_rows),
        "",
        "## 13. 与普通全文检索的对比",
        "",
        "A 方案只检索原始文本 chunk，优点是直接、透明，适合作为基线。B 方案使用专业词汇库、AST 节点和 Wiki 页面，能够返回专业术语、AST 节点、Wiki 页面和来源文献，因此解释性和可追溯性更强。B 方案的不足是会继承规则抽取产生的噪声，尤其当 chunk 来自目录页或长条款时，候选术语可能过长或过宽。",
        "",
        "## 14. 当前方案优势",
        "",
        "- 保留从原始文献到文本、chunk、AST 节点、Wiki 页面和检索结果的链路。",
        "- 将法律主体、行为、对象、条件、义务、后果等字段显式化，便于后续检索和审核。",
        "- Markdown Wiki 适合 Obsidian、人类维护和 LLM/Agent 调用。",
        "- 在没有外部 LLM API 的情况下，仍可用本地规则和 TF-IDF 完成可运行原型。",
        "",
        "## 15. 当前方案不足",
        "",
        "- 当前九篇文献只是原型验证样本，不代表完整法律知识库。",
        "- 规则解析无法替代专业法律解释，可能产生过度抽取、漏抽取或字段分类错误。",
        "- 词汇库、口语映射和 AST 节点均为候选结果，需要人工审核。",
        "- 检索验证没有人工标注的 gold set，因此只能说明原型可运行，不能证明法律准确性。",
        "",
        "## 16. 后续优化方向",
        "",
        "- 建立人工审核工作流，为 AST 节点、Wiki 页面和词汇条目添加 approved/rejected 状态。",
        "- 引入更细的文献类型识别、法律领域过滤和条文层级识别。",
        "- 使用人工标注问题集评估 Top-K 相关性、召回率、准确率和证据质量。",
        "- 在环境变量配置 API Key 的前提下，可选接入 LLM 或 embedding 模型辅助解析和检索。",
        "- 加入增量构建能力，避免每次运行都重建全部 Wiki 页面。",
        "",
        "## 17. 结论",
        "",
        "本原型证明“类 AST + LLM Wiki”用于法律/政策文献知识组织是可行的。它不能直接给出权威法律结论，但能把原始文献转化为带来源证据的结构化候选知识，并支持专业术语映射、Wiki 组织和可解释检索。当前结果适合作为后续 Agent/LLM 问答、人工审核和检索增强实验的基础。",
        "",
        "> 重要提示：本报告和所有生成内容仅为原型验证材料，不构成法律意见。所有候选 AST 节点、Wiki 页面、词汇库和检索结果均应保持 `review_status = pending`，直到经过人工审核。",
    ]

    return "\n".join(lines) + "\n"


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell: Any, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)


def add_markdown_table(document: Document, table_lines: list[str]) -> None:
    rows = []
    for line in table_lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return

    table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for col_index in range(len(table.columns)):
            cell = table.cell(row_index, col_index)
            text = row[col_index] if col_index < len(row) else ""
            set_cell_text(cell, text, bold=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
    document.add_paragraph()


def add_markdown_to_docx(document: Document, markdown_text: str) -> None:
    lines = markdown_text.splitlines()
    table_buffer: list[str] = []

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(document, table_buffer)
            table_buffer = []

    for line in lines:
        if line.strip().startswith("|"):
            table_buffer.append(line)
            continue
        flush_table()

        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            paragraph = document.add_heading(stripped[2:], level=0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=2)
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped.startswith("> "):
            paragraph = document.add_paragraph(stripped[2:])
            paragraph.style = document.styles["Intense Quote"]
        else:
            document.add_paragraph(stripped)
    flush_table()


def style_docx(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    for style_name in ["Normal", "Body Text"]:
        if style_name in document.styles:
            style = document.styles[style_name]
            style.font.name = "Microsoft YaHei"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            style.font.size = Pt(10.5)

    for style_name, size, color in [
        ("Title", 18, "17365D"),
        ("Heading 1", 14, "17365D"),
        ("Heading 2", 12, "1F4E79"),
    ]:
        if style_name in document.styles:
            style = document.styles[style_name]
            style.font.name = "Microsoft YaHei"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15


def write_docx_report(markdown_text: str, docx_path: Path) -> None:
    document = Document()
    style_docx(document)
    add_markdown_to_docx(document, markdown_text)
    style_docx(document)
    backup_existing_file(docx_path)
    document.save(docx_path)


def generate_final_report(
    *,
    project_root: Path = PROJECT_ROOT,
    output_dir: Path | None = None,
    log=print,
) -> dict[str, Any]:
    output_dir = output_dir or project_root / "outputs/reports"
    output_dir.mkdir(parents=True, exist_ok=True)
    markdown_path = output_dir / FINAL_REPORT_MD_NAME
    docx_path = output_dir / FINAL_REPORT_DOCX_NAME

    metrics = collect_project_metrics(project_root)
    markdown_text = build_final_report_markdown(metrics)

    backup_existing_file(markdown_path)
    markdown_path.write_text(markdown_text, encoding="utf-8-sig")
    write_docx_report(markdown_text, docx_path)

    log(f"Final Markdown report saved: {markdown_path}")
    log(f"Final DOCX report saved: {docx_path}")
    return {
        "markdown_path": markdown_path,
        "docx_path": docx_path,
        "metrics": metrics,
    }
